"""
Re-derive the 'Gain Error' column of an RX phase LUT xlsx from a gain/phase sweep taken with that LUT
(csv written by tests/char/rx_Av_iq_sweep.py with sweep = 'both').

Per gain code the gain across the phase states is centered on its median, then averaged over the gain
codes. A positive code lowers the gain by 0.5 dB (verified: sweeps with and without the codes differ by
about -0.5 dB per code), so the gain deviation that must be cancelled is
    t(state) = measured deviation + 0.5 * code used in the sweep
and the new code is the 3-bit two's complement value (-4...3) that leaves the flattest gain.
The LUT is updated in place after a backup copy is made (earlier backups are kept).

exclude_g gain codes are left out of the derivation (their predicted effect is still printed).
lut_used is the LUT the sweep was actually taken with, when it differs from the file being updated.

usage: python gain_err_from_sweep.py
"""
import os
import shutil
import numpy as np
import pandas as pd
import openpyxl

sweep_csv = r'..\tests\char\logs\rx_Av_iq_sweep__v2__ant_sel_1__bias_NOM__sweep_both__g_[30,10,60]__p_[4,1,124]__freq_9p5__2026-10-03_16-05-17.csv'
lut_xlsx = r'logs\v2__rx_phase_lut__freq_9p5__gm_1__backoff_5__algo_joint_max__av_2047__bias_nom__vdd_3p3__temp_25C.xlsx'
exclude_g = (60,)   # gain codes left out when deriving the correction
lut_used = lut_xlsx.replace('.xlsx', '__before_av_gain_corr_2.xlsx')   # LUT the sweep was taken with (None = lut_xlsx)

states, pad = 121, 4          # entries pad...pad+states-1 are the phase states, the others wrap around
glitch_db = 0.5               # first point after a gain code change is dropped if it is this far from the next one

# ---------------- measured gain deviation per state ----------------
d = pd.read_csv(sweep_csv)
G = d.pivot(index='p_idx', columns='g_idx', values='Gain dB')          # [p_idx 4...124, gain code]
G = G.where(~((G.index == G.index[0])[:, None] & ((G.iloc[0] - G.iloc[1]).abs() > glitch_db).values[None, :]))   # settling glitch
Y = G - G.median()                                                      # deviation per gain code
y = Y.drop(columns=list(exclude_g), errors='ignore').mean(axis=1).values   # averaged over the used gain codes

# ---------------- codes used in the sweep -> new codes ----------------
ws_used = openpyxl.load_workbook(lut_used or lut_xlsx)['Sheet1']
old = np.array([ws_used.cell(row=n + 2, column=3).value for n in range(pad, pad + states)])
wb = openpyxl.load_workbook(lut_xlsx)
ws = wb['Sheet1']
t = y + 0.5 * old                                                       # deviation without any correction

best = None
for offset in np.arange(-3, 3, 0.05):                                   # level shift that keeps the codes in range
    code = np.clip(np.round((t - offset) / 0.5), -4, 3).astype(int)
    score = np.std(t - 0.5 * code)
    if best is None or score < best[0]:
        best = (score, code)
new = best[1]

# ---------------- report ----------------
res = Y.values + 0.5 * old[:, None] - 0.5 * new[:, None]                # predicted deviation per gain code after the change
print(f'gain codes used: {[g for g in Y.columns if g not in exclude_g]}')
print(f'{"gain code":>9}{"sd before":>11}{"sd after":>10}{"p-p before":>12}{"p-p after":>11}   (dB, over the {states} states)')
for j, g in enumerate(Y.columns):
    b, a = Y[g].dropna(), pd.Series(res[:, j]).dropna()
    print(f'{g:>9}{b.std():11.2f}{a.std():10.2f}{np.ptp(b):12.2f}{np.ptp(a):11.2f}')
print(f'codes: old {old.min()}...{old.max()}, new {new.min()}...{new.max()}, changed {(new != old).sum()} of {states}')

# ---------------- write ----------------
n_backup = 1
backup = lut_xlsx.replace('.xlsx', '__before_av_gain_corr.xlsx')
while os.path.exists(backup):   # never overwrite an earlier backup
    n_backup += 1
    backup = lut_xlsx.replace('.xlsx', f'__before_av_gain_corr_{n_backup}.xlsx')
shutil.copy(lut_xlsx, backup)
for n in range(128):
    ws.cell(row=n + 2, column=3).value = int(new[(n - pad) % states])
wb.save(lut_xlsx)
print(f'\nUpdated {os.path.abspath(lut_xlsx)}\nBackup  {os.path.abspath(backup)}')
